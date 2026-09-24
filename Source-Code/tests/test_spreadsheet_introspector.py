"""Tests for tgdatabridge.core.spreadsheet_introspector -- turning a worksheet
into a Schema.

Most cases use CSV (no third-party package needed); the few that genuinely
need a workbook guard themselves with a plain ImportError check, the same
way tests/test_spreadsheet_connector.py does.
"""
import pathlib
import tempfile

from tgdatabridge.core.spreadsheet_introspector import introspect_schema
from tgdatabridge.db.base import ConnectionParams
from tgdatabridge.db.spreadsheet_connector import SpreadsheetConnector


def _connected(text, name="data.csv", schema=None):
    path = pathlib.Path(tempfile.mkdtemp(prefix="tgdatabridge_sheet_")) / name
    path.write_text(text, encoding="utf-8")
    conn = SpreadsheetConnector(ConnectionParams(
        host="", port=0, database=str(path), username="", password="", schema=schema))
    conn.connect()
    return conn


def _schema_from(text, name="data.csv"):
    conn = _connected(text, name)
    return introspect_schema(conn, conn.schema_name)


def _column(table, name):
    for column in table.columns:
        if column.name == name:
            return column
    raise AssertionError(f"no column {name!r} in {[c.name for c in table.columns]}")


def _messages(issues):
    return " | ".join(i.message for i in issues)


# ------------------------------------------------------------ basics

def test_source_engine_is_recorded_on_the_schema():
    schema = _schema_from("id\n1\n")
    assert schema.source_engine == "Excel/CSV"


def test_one_sheet_becomes_one_table():
    schema = _schema_from("id,name\n1,Alice\n")
    assert len(schema.tables) == 1
    assert schema.tables[0].name == "data"


def test_header_row_becomes_columns():
    table = _schema_from("id,name,email\n1,Alice,a@x.com\n").tables[0]
    assert [c.name for c in table.columns] == ["id", "name", "email"]


def test_row_count_estimate_counts_data_rows_only():
    table = _schema_from("id\n1\n2\n3\n").tables[0]
    assert table.row_count_estimate == 3


def test_source_collection_records_the_real_sheet_name():
    # Migration looks the sheet up through this, since the table name is
    # the *sanitized* form.
    table = _schema_from("id\n1\n", name="Q1 Sales.csv").tables[0]
    assert table.name == "Q1_Sales"
    assert table.source_collection == "Q1 Sales"


def test_nothing_relational_is_invented():
    # A worksheet has no keys, indexes, views, sequences or routines, and
    # guessing at a PRIMARY KEY from a column that merely looks like an ID
    # would break Apply DDL the moment a duplicate showed up.
    schema = _schema_from("id,name\n1,Alice\n2,Bob\n")
    table = schema.tables[0]
    assert table.constraints == []
    assert table.indexes == []
    assert schema.views == []
    assert schema.sequences == []
    assert schema.routines == []


# -------------------------------------------------------- type inference

def test_integer_column_gets_a_native_integer_type():
    # NUMBER(9)/NUMBER(18) map to a real INTEGER/BIGINT on every target;
    # NUMBER(19) would silently degrade to DECIMAL everywhere.
    table = _schema_from("id\n1\n2\n3\n").tables[0]
    assert _column(table, "id").data_type == "NUMBER(9)"


def test_wide_integer_column_widens_to_bigint():
    table = _schema_from("id\n1234567890123\n").tables[0]
    assert _column(table, "id").data_type == "NUMBER(18)"


def test_decimal_column_becomes_a_float_type():
    table = _schema_from("amount\n1.5\n2.25\n").tables[0]
    assert _column(table, "amount").data_type == "BINARY_DOUBLE"


def test_int_and_float_mix_widens_without_a_warning():
    # 100.50 / 0 / -42.75 is just a money column.
    table = _schema_from("amount\n100.50\n0\n-42.75\n").tables[0]
    column = _column(table, "amount")
    assert column.data_type == "BINARY_DOUBLE"
    assert not [i for i in column.source_issues if i.severity == "warning"]


def test_text_column_is_sized_from_the_longest_sampled_value():
    table = _schema_from("name\nAl\nAlexandra\n").tables[0]
    data_type = _column(table, "name").data_type
    assert data_type.startswith("VARCHAR2(")
    # Padded above the observed maximum, since only a sample was read.
    assert int(data_type[len("VARCHAR2("):-1]) > len("Alexandra")


def test_iso_date_column_is_recognized():
    table = _schema_from("when\n2024-01-15\n2024-02-20\n").tables[0]
    assert _column(table, "when").data_type == "DATE"


def test_boolean_column_is_recognized():
    table = _schema_from("flag\nTRUE\nfalse\n").tables[0]
    assert _column(table, "flag").data_type == "BOOLEAN"


def test_leading_zero_codes_stay_text():
    table = _schema_from("zip\n01234\n02115\n").tables[0]
    assert _column(table, "zip").data_type.startswith("VARCHAR2(")


def test_genuinely_mixed_column_becomes_text_and_is_flagged():
    table = _schema_from("amount\n1\n2\nN/A\n").tables[0]
    column = _column(table, "amount")
    assert column.data_type.startswith("VARCHAR2(")
    warnings = [i for i in column.source_issues if i.severity == "warning"]
    assert len(warnings) == 1
    assert "mixes values of different kinds" in warnings[0].message


def test_all_empty_column_falls_back_to_wide_text_with_a_warning():
    table = _schema_from("a,b\n1,\n2,\n").tables[0]
    column = _column(table, "b")
    assert column.data_type == "VARCHAR2(4000)"
    assert [i for i in column.source_issues if i.severity == "warning"]


def test_every_column_is_nullable_even_with_no_blanks_in_the_sample():
    # NOT NULL could only ever be guessed from a sample, and that guess
    # fails in the worst way: the migration runs for an hour and then
    # aborts on the first blank cell the full file turns out to have.
    table = _schema_from("id,name\n1,Alice\n2,Bob\n").tables[0]
    assert all(c.nullable for c in table.columns)


# ------------------------------------------------------------- naming

def test_illegal_sheet_name_is_sanitized_and_recorded():
    schema = _schema_from("id\n1\n", name="Order Items (2024).csv")
    table = schema.tables[0]
    assert table.name == "Order_Items_2024"
    assert "isn't a legal SQL identifier" in _messages(table.issues)


def test_illegal_header_is_sanitized_and_recorded():
    table = _schema_from("Unit Price $\n1\n").tables[0]
    assert [c.name for c in table.columns] == ["Unit_Price"]
    assert "Unit Price $" in _messages(table.issues)


def test_blank_header_cell_gets_a_positional_name():
    table = _schema_from("id,,name\n1,x,Alice\n").tables[0]
    assert [c.name for c in table.columns] == ["id", "column_2", "name"]
    assert "blank" in _messages(table.issues)


def test_trailing_blank_headers_are_dropped_not_named():
    # A used range wider than the header is normal and shouldn't produce
    # phantom columns.
    table = _schema_from("id,name,,\n1,Alice,,\n").tables[0]
    assert [c.name for c in table.columns] == ["id", "name"]


def test_duplicate_headers_are_deduplicated():
    table = _schema_from("id,id\n1,2\n").tables[0]
    assert [c.name for c in table.columns] == ["id", "id_2"]
    assert "duplicate" in _messages(table.issues).lower()


def test_duplicate_detection_is_case_insensitive():
    # Most targets fold or compare identifiers case-insensitively, so
    # "Total" and "total" would collide there even though Python sees two
    # different strings.
    table = _schema_from("Total,total\n1,2\n").tables[0]
    assert [c.name for c in table.columns] == ["Total", "total_2"]


def test_headers_differing_only_in_illegal_characters_are_deduplicated():
    table = _schema_from("unit price,unit-price\n1,2\n").tables[0]
    assert [c.name for c in table.columns] == ["unit_price", "unit_price_2"]


# -------------------------------------------------------- malformed input

def test_header_only_sheet_produces_a_table_with_a_warning():
    schema = _schema_from("id,name\n")
    table = schema.tables[0]
    assert [c.name for c in table.columns] == ["id", "name"]
    assert "no data rows" in _messages(table.issues)


def test_completely_empty_file_produces_no_tables():
    schema = _schema_from("")
    assert schema.tables == []


def test_sheet_with_a_blank_first_row_is_reported_not_guessed():
    # Rather than hunting for a header further down -- that heuristic
    # fails silently and confusingly when it guesses wrong.
    schema = _schema_from(",,\nid,name\n1,Alice\n")
    table = schema.tables[0]
    assert table.columns == []
    assert "completely empty first row" in _messages(table.issues)


def test_ragged_rows_do_not_break_inference():
    schema = _schema_from("a,b,c\n1\n2,3,4\n")
    table = schema.tables[0]
    assert [c.name for c in table.columns] == ["a", "b", "c"]


# ----------------------------------------------------------- workbooks

def test_every_sheet_in_a_workbook_becomes_a_table():
    try:
        import openpyxl
    except ImportError:
        return
    path = pathlib.Path(tempfile.mkdtemp(prefix="tgdatabridge_sheet_")) / "book.xlsx"
    workbook = openpyxl.Workbook()
    workbook.remove(workbook.active)
    first = workbook.create_sheet("Customers")
    first.append(["id", "name"])
    first.append([1, "Alice"])
    second = workbook.create_sheet("Orders")
    second.append(["order_id"])
    second.append([99])
    workbook.create_sheet("Blank")
    workbook.save(path)

    conn = SpreadsheetConnector(ConnectionParams(
        host="", port=0, database=str(path), username="", password="", schema=None))
    conn.connect()
    schema = introspect_schema(conn, conn.schema_name)
    conn.close()

    # The genuinely empty sheet contributes nothing rather than an
    # unusable zero-column table.
    assert [t.name for t in schema.tables] == ["Customers", "Orders"]
    assert [c.name for c in schema.tables[0].columns] == ["id", "name"]
